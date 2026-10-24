"""Knowledge document upload (POST /api/knowledge/upload) — .pdf/.docx/.txt →
KB article, with validation (type/size/empty), text normalization and role guard.
Runs offline: OpenAI is disabled by conftest, so the auto-reindex embeds nothing.
"""
import io

from app.routers.knowledge import BODY_MAX_CHARS


def _upload(client, name, data, content_type="application/octet-stream"):
    return client.post("/api/knowledge/upload",
                       files={"file": (name, io.BytesIO(data), content_type)})


def _make_docx(paragraphs, table_rows=()):
    import docx
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, cell in enumerate(row):
                table.rows[r].cells[c].text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pdf(text):
    """Hand-build a minimal one-page PDF whose content stream pypdf can extract."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_at}\n%%EOF").encode()
    return out


def test_upload_txt_creates_article(auth_client):
    body = "Opening hours\n\n\n\n\nMon-Fri 10:00-19:00   \nWalk-ins welcome\x00"
    r = _upload(auth_client, "Clinic Hours.txt", body.encode("utf-8"))
    assert r.status_code == 200, r.text
    a = r.json()
    assert a["title"] == "Clinic Hours"
    assert a["type"] == "general" and a["active"] is True
    # Normalized: control chars stripped, trailing spaces trimmed, blank runs collapsed.
    assert a["content"] == "Opening hours\n\nMon-Fri 10:00-19:00\nWalk-ins welcome"
    # And it is listed with the clinic's articles.
    titles = [x["title"] for x in auth_client.get("/api/knowledge/articles").json()]
    assert "Clinic Hours" in titles


def test_upload_txt_thai_cp874_fallback(auth_client):
    r = _upload(auth_client, "thai.txt", "โปรโมชันเลเซอร์".encode("cp874"))
    assert r.status_code == 200, r.text
    assert r.json()["content"] == "โปรโมชันเลเซอร์"


def test_upload_docx_paragraphs_and_tables(auth_client):
    data = _make_docx(["Botox aftercare", "Avoid massage for 24h"],
                      table_rows=[["Area", "Price"], ["Forehead", "4900"]])
    r = _upload(auth_client, "aftercare.docx", data)
    assert r.status_code == 200, r.text
    content = r.json()["content"]
    assert "Botox aftercare" in content and "Avoid massage for 24h" in content
    assert "Forehead | 4900" in content


def test_upload_pdf(auth_client):
    r = _upload(auth_client, "policy.pdf", _make_pdf("Deposit is refundable within 7 days"))
    assert r.status_code == 200, r.text
    assert "Deposit is refundable within 7 days" in r.json()["content"]


def test_upload_corrupt_pdf_400(auth_client):
    r = _upload(auth_client, "broken.pdf", b"%PDF-1.4 garbage not a real pdf")
    assert r.status_code == 400
    assert "PDF" in r.json()["detail"]


def test_upload_rejects_unsupported_extension(auth_client):
    r = _upload(auth_client, "photo.png", b"\x89PNG....")
    assert r.status_code == 400


def test_upload_rejects_empty_file(auth_client):
    r = _upload(auth_client, "empty.txt", b"")
    assert r.status_code == 400


def test_upload_rejects_oversized_file(auth_client):
    r = _upload(auth_client, "big.txt", b"x" * (10 * 1024 * 1024 + 1))
    assert r.status_code == 400
    assert "10 MB" in r.json()["detail"]


def test_upload_body_is_capped_for_embedding(auth_client):
    r = _upload(auth_client, "huge.txt", b"a" * (BODY_MAX_CHARS + 5000))
    assert r.status_code == 200, r.text
    assert len(r.json()["content"]) == BODY_MAX_CHARS


def test_upload_requires_editor_role(client, clinic, make_owner):
    _u, pw = make_owner(clinic, email="sales@a.co", role="sales")
    r = client.post("/api/auth/login", json={"email": "sales@a.co", "password": pw})
    assert r.status_code == 200, r.text
    r = client.post(f"/api/clinics/{clinic.id}/enter")
    assert r.status_code == 200, r.text
    r = _upload(client, "notes.txt", b"hello")
    assert r.status_code == 403
